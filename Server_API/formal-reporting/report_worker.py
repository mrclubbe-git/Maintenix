#!/usr/bin/env python3
"""Deterministic DOCX/PDF worker for Maintenix formal reporting."""

from __future__ import annotations

import copy
import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Mm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


MAINTENIX_ROOT = Path(__file__).resolve().parents[1]
TEMPLATES_DIR = MAINTENIX_ROOT / "data" / "templates" / "formal-reports"
DOCX_DIR = MAINTENIX_ROOT / "data" / "uploads" / "reports"
PDF_DIR = MAINTENIX_ROOT / "data" / "uploads" / "reports_pdf"
EXPECTED_TEMPLATES = {
    "CRA_Conveyor_Fire_Service_Report_Master_Template.docx",
    "CRA_Server_Room_Fire_Service_Report_Master_Template.docx",
    "CRA_Substation_Fire_Service_Report_Master_Template.docx",
}
MODEL_CONFIG = {
    "Substation": {
        "template": "CRA_Substation_Fire_Service_Report_Master_Template.docx",
        "category": "Substation Fire System",
    },
    "Conveyor": {
        "template": "CRA_Conveyor_Fire_Service_Report_Master_Template.docx",
        "category": "Conveyor Fire System",
    },
    "Server Room": {
        "template": "CRA_Server_Room_Fire_Service_Report_Master_Template.docx",
        "category": "Server Room Fire System",
    },
}


def inspect_template(path: Path) -> dict:
    document = Document(path)
    headings = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    return {"name": path.name, "size": path.stat().st_size, "tables": len(document.tables),
            "paragraphs": len(document.paragraphs), "text_preview": headings[:8]}


def health() -> int:
    present = {path.name for path in TEMPLATES_DIR.glob("*.docx")}
    templates = [inspect_template(path) for path in sorted(TEMPLATES_DIR.glob("*.docx"))]
    result = {"ok": present == EXPECTED_TEMPLATES and shutil.which("libreoffice") is not None,
              "python_docx": True, "libreoffice": shutil.which("libreoffice"),
              "template_count": len(templates), "missing_templates": sorted(EXPECTED_TEMPLATES - present),
              "unexpected_templates": sorted(present - EXPECTED_TEMPLATES), "templates": templates}
    print(json.dumps(result))
    return 0 if result["ok"] else 1


def report_model(section: str, system_name: str) -> str:
    name = system_name.lower()
    if section == "IIT & IT" or any(word in name for word in ("server room", "ups room", "archive", "control room")):
        return "Server Room"
    if "conveyor" in name:
        return "Conveyor"
    return "Substation"


def safe_slug(value: str) -> str:
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", value.lower())).strip("-")


def set_cell(cell, value: str, color: str | None = None) -> None:
    paragraph = cell.paragraphs[0]
    seed = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    font_name, font_size, bold = seed.font.name, seed.font.size, seed.bold
    alignment = paragraph.alignment
    for extra_paragraph in cell.paragraphs[1:]:
        extra_paragraph._element.getparent().remove(extra_paragraph._element)
    paragraph.clear()
    paragraph.alignment = alignment
    run = paragraph.add_run(str(value))
    run.font.name, run.font.size, run.bold = font_name, font_size, bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def clone_data_rows(table, count: int) -> list:
    prototype = table.rows[1]._tr
    table._tbl.remove(prototype)
    rows = []
    for _ in range(count):
        table._tbl.append(copy.deepcopy(prototype))
        row = table.rows[-1]
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))
        rows.append(row)
    return rows


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def align_tables_to_first(document) -> None:
    """Give every table the first table's outer left/right boundaries."""
    if not document.tables:
        return
    first_grid = [int(col.get(qn("w:w")) or 0) for col in document.tables[0]._tbl.tblGrid.gridCol_lst]
    guide_width = sum(first_grid)
    if guide_width <= 0:
        return

    for table in document.tables:
        grid_columns = table._tbl.tblGrid.gridCol_lst
        widths = [int(col.get(qn("w:w")) or 0) for col in grid_columns]
        current_width = sum(widths)
        if current_width <= 0:
            continue
        scale = guide_width / current_width
        scaled = [max(1, round(width * scale)) for width in widths]
        scaled[-1] += guide_width - sum(scaled)
        for column, width in zip(grid_columns, scaled):
            column.set(qn("w:w"), str(width))

        for cell_width in table._tbl.xpath(".//w:tcW"):
            if cell_width.get(qn("w:type"), "dxa") != "dxa":
                continue
            value = int(cell_width.get(qn("w:w")) or 0)
            if value > 0:
                cell_width.set(qn("w:w"), str(max(1, round(value * scale))))

        properties = table._tbl.tblPr
        table_width = properties.find(qn("w:tblW"))
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            properties.insert(0, table_width)
        table_width.set(qn("w:w"), str(guide_width))
        table_width.set(qn("w:type"), "dxa")

        layout = properties.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            properties.append(layout)
        layout.set(qn("w:type"), "fixed")
        table.alignment = document.tables[0].alignment
        table.autofit = False


def replace_paragraph_tokens(document, replacements: dict[str, str]) -> None:
    for paragraph in document.paragraphs:
        original = paragraph.text
        updated = original
        for token, value in replacements.items():
            updated = updated.replace(token, value)
        if updated != original:
            paragraph.text = updated


def build_report(payload: dict, model: str, systems: list[dict]) -> dict:
    config = MODEL_CONFIG[model]
    document = Document(TEMPLATES_DIR / config["template"])
    for doc_section in document.sections:
        doc_section.page_width = Mm(210)
        doc_section.page_height = Mm(297)
    section, month, week = payload["section"], payload["month"], int(payload["week"])
    year, month_number = (int(part) for part in month.split("-"))
    month_name = date(year, month_number, 1).strftime("%B %Y")
    replace_paragraph_tokens(document, {"[SECTION / AREA]": section, "[MONTH YEAR]": month_name.upper()})

    metadata = document.tables[0]
    set_cell(metadata.cell(0, 1), f"{month_name} — Week {week}")
    set_cell(metadata.cell(0, 3), payload.get("weekRange", ""))
    set_cell(metadata.cell(1, 1), section)
    set_cell(metadata.cell(1, 3), "Routine service")
    set_cell(metadata.cell(2, 3), "Final")

    serviced_rows = clone_data_rows(document.tables[1], len(systems))
    finding_rows = clone_data_rows(document.tables[2], len(systems))
    defect_systems = 0
    for index, (entry, service_row, finding_row) in enumerate(zip(systems, serviced_rows, finding_rows), 1):
        system, response = entry["system"], entry["response"]
        has_defects = response["defectsFound"] == "yes"
        defect_systems += int(has_defects)
        status = "Defect Identified" if has_defects else "No Defects"
        service_text = "Service completed; defect(s) identified." if has_defects else "Review completed; no defects identified."
        findings = response.get("defects", []) if has_defects else []
        finding_text = "\n".join(f"{i}. {item['finding'].strip()}" for i, item in enumerate(findings, 1)) or "1. No defects identified."
        action_text = "\n".join(f"{i}. {item['action'].strip()}" for i, item in enumerate(findings, 1)) or "No action required."
        values = [index, system["name"], config["category"], "Monthly", service_text, status]
        for cell, value in zip(service_row.cells, values): set_cell(cell, value)
        set_cell(service_row.cells[5], status, "C00000" if has_defects else "008000")
        for cell, value in zip(finding_row.cells, [system["name"], status, finding_text, action_text]): set_cell(cell, value)

    feedback = document.tables[3]
    set_cell(feedback.cell(1, 1), f"Week {week} service review completed for {len(systems)} listed system(s).")
    set_cell(feedback.cell(1, 2), "Review the recorded findings and actions.")
    set_cell(feedback.cell(2, 1), f"{defect_systems} system(s) with defects; {len(systems) - defect_systems} with no defects.")
    set_cell(feedback.cell(2, 2), "Complete the required corrective actions." if defect_systems else "No action required.")
    set_cell(feedback.cell(3, 1), "Continue routine monitoring in accordance with the service programme.")
    set_cell(feedback.cell(3, 2), "No additional action required." if not defect_systems else "Track open corrective actions to completion.")

    compiled = document.tables[4]
    set_cell(compiled.cell(1, 0), payload["compiledName"].strip())
    set_cell(compiled.cell(1, 1), payload["compiledDesignation"].strip())
    set_cell(compiled.cell(1, 2), date.today().strftime("%d %B %Y"))
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("5. Compiled By"):
            paragraph.paragraph_format.keep_with_next = True
    for row_index, row in enumerate(compiled.rows):
        prevent_row_split(row)
        if row_index == 0:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True

    align_tables_to_first(document)

    generation_id = safe_slug(str(payload.get("generationId", "")))
    suffix = f"-{generation_id}" if generation_id else ""
    base = f"{safe_slug(section)}-{safe_slug(model)}-{month}-week-{week}{suffix}"
    docx_path, pdf_path = DOCX_DIR / f"{base}.docx", PDF_DIR / f"{base}.pdf"
    DOCX_DIR.mkdir(parents=True, exist_ok=True); PDF_DIR.mkdir(parents=True, exist_ok=True)
    document.save(docx_path)
    with tempfile.TemporaryDirectory(prefix="service-report-") as temp_dir:
        profile_dir = Path(temp_dir) / "libreoffice-profile"
        profile_dir.mkdir()
        process = subprocess.run(["libreoffice", f"-env:UserInstallation={profile_dir.as_uri()}", "--headless",
                                  "--convert-to", "pdf", "--outdir", temp_dir, str(docx_path)],
                                 capture_output=True, text=True, timeout=120)
        converted = Path(temp_dir) / f"{docx_path.stem}.pdf"
        if process.returncode != 0 or not converted.exists():
            raise RuntimeError(process.stderr.strip() or process.stdout.strip() or "PDF conversion failed")
        shutil.copy2(converted, pdf_path)
    text_result = subprocess.run(["pdftotext", "-layout", str(pdf_path), "-"], capture_output=True, text=True, timeout=30)
    if text_result.returncode != 0:
        raise RuntimeError("Generated PDF could not be verified")
    pdf_text = text_result.stdout
    forbidden = ("[ADDITIONAL FINDING", "[DETAILED FINDING]", "[PROTECTED AREA / SYSTEM]")
    if any(token in pdf_text for token in forbidden):
        raise RuntimeError("Generated PDF contains an unresolved template placeholder")
    pages = pdf_text.split("\f")
    heading_page = next((index for index, page in enumerate(pages) if "5. Compiled By" in page), None)
    values_page = next((index for index, page in enumerate(pages) if payload["compiledName"].strip() in page), None)
    if heading_page is None or values_page is None or heading_page != values_page:
        raise RuntimeError("Compiled By section split across pages")
    return {"model": model, "systems": len(systems), "docx": docx_path.name, "pdf": pdf_path.name}


def generate() -> int:
    payload = json.load(sys.stdin)
    responses = {str(item["systemId"]): item for item in payload["responses"]}
    groups: dict[str, list] = {}
    for system in payload["systems"]:
        response = responses.get(str(system["id"]))
        if not response:
            raise ValueError(f"Missing response for system {system['id']}")
        model = report_model(payload["section"], system["name"])
        groups.setdefault(model, []).append({"system": system, "response": response})
    outputs = [build_report(payload, model, systems) for model, systems in groups.items()]
    print(json.dumps({"ok": True, "outputs": outputs}))
    return 0


def main() -> int:
    command = sys.argv[1] if len(sys.argv) > 1 else "health"
    if command == "health": return health()
    if command == "generate": return generate()
    print(json.dumps({"ok": False, "error": f"Unknown command: {command}"}))
    return 2


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(json.dumps({"ok": False, "error": str(error)}))
        raise SystemExit(1)
